"""
bau_deck.py — baut das Foliendeck "Fallstudie mini-contacts".

Voraussetzungen:
  * THWS-Skill unter ~/.claude/skills/thws-slides (Template, Bausteine, Spec)
  * /opt/anaconda3/bin/python3 (python-pptx, lxml)

Aufruf:
    /opt/anaconda3/bin/python3 slides/bau_deck.py
Ergebnis:
    slides/Fallstudie_mini_contacts.pptx
"""
import re
import sys
from pathlib import Path

SKILL = Path.home() / ".claude" / "skills" / "thws-slides"
sys.path.insert(0, str(SKILL / "scripts"))

import spec_loader                                        # noqa: E402
spec_loader.SPECS_DIR = SKILL / "assets" / "specs"        # Specs liegen unter assets/
import bausteine as B                                     # noqa: E402
from bausteine import (E, kachel, band, chevron_kette, gegenueber,      # noqa: E402
                       zuordnung, kennzahlen, sprechblase, shape, label,
                       textbox, kopf, rgb, BLUE, GOOD, WARN, METHOD, HINT,
                       BODY, SEC, SAND, BAND, WHITE, HAIR, CL, CR, CW,
                       C3, W3, C2, W2)
from pptx import Presentation                              # noqa: E402
from pptx.enum.shapes import MSO_SHAPE                     # noqa: E402
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN            # noqa: E402
from pptx.util import Pt                                   # noqa: E402

HIER = Path(__file__).parent
ZIEL = HIER / "Fallstudie_mini_contacts.pptx"
IMG = HIER.parent / "docs" / "img"

KEYWORD = "FF40FF"
SQL_KEYWORDS = {
    "CREATE", "TABLE", "INTEGER", "PRIMARY", "KEY", "NOT", "NULL",
    "REFERENCES", "ON", "DELETE", "CASCADE", "TEXT", "CHECK", "IN",
    "PRAGMA", "INSERT", "SELECT", "UPDATE", "FROM", "WHERE", "UNIQUE",
}

prs = Presentation(str(SKILL / "assets" / "template.pptx"))
master = prs.slide_masters[0]
LAYOUT = {l.name: l for l in master.slide_layouts}

# Fusszeile: Datum auf den Stand dieses Decks setzen (Kurs stimmt bereits).
# Steht das Datum in einem Feld (keine Runs), bleibt es unangetastet.
for sh in master.shapes:
    if sh.has_text_frame and sh.text_frame.text.strip() == "25.08.2026":
        for p in sh.text_frame.paragraphs:
            for r in p.runs:
                if r.text.strip() == "25.08.2026":
                    r.text = "26.08.2026"


# ------------------------------------------------------------------ Helfer
def folie(layout_name):
    return prs.slides.add_slide(LAYOUT[layout_name])


def einl(slide, text, variante="Lisa_Slide"):
    """Einleitung in Platzhalter idx 14, 14 pt, Geometrie vollstaendig."""
    ph = B.einleitung(slide, text, variante)
    vorlage = [q for q in slide.slide_layout.placeholders
               if q.placeholder_format.idx == 14][0]
    ph.left, ph.top, ph.width = vorlage.left, vorlage.top, vorlage.width
    ph.height = Pt(63.8)                       # traegt vier Zeilen bei 14 pt
    for p in ph.text_frame.paragraphs:
        p.line_spacing = 0.95
        for r in p.runs:
            r.font.size = Pt(14)
    return ph


def mono_box(slide, x, y, w, h, zeilen, size=13, color=BODY):
    """Codeblock in Consolas; SQL-Schluesselwoerter im Bestandston."""
    tb = slide.shapes.add_textbox(E(x), E(y), E(w), E(h))
    tf = tb.text_frame
    tf.word_wrap = False
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    for i, zeile in enumerate(zeilen):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.line_spacing = 1.0
        for tok in re.split(r"(\w+)", zeile):
            if not tok:
                continue
            r = p.add_run()
            r.text = tok
            r.font.name = "Consolas"
            r.font.size = Pt(size)
            r.font.color.rgb = rgb(KEYWORD if tok in SQL_KEYWORDS else color)
    return tb


def hline(slide, x, y, w):
    shape(slide, MSO_SHAPE.RECTANGLE, x, y - 0.5, w, 1.0, BODY)


def vline(slide, x, y, h):
    shape(slide, MSO_SHAPE.RECTANGLE, x - 0.5, y, 1.0, h, BODY)


def kardlabel(slide, x, y, txt):
    textbox(slide, x, y, 22, 15, [(txt, True, BLUE)], 12,
            align=PP_ALIGN.CENTER)


def tabellenbox(slide, x, y, w, name, zeilen):
    """Kleine Schemabox: blauer Kopf, eine Zeile je Spalte."""
    kopf_h, zeilen_h = 22.0, 15.0
    h = kopf_h + zeilen_h * len(zeilen)
    shape(slide, MSO_SHAPE.RECTANGLE, x, y, w, h, WHITE, HAIR)
    sh = shape(slide, MSO_SHAPE.RECTANGLE, x, y, w, kopf_h, BLUE)
    label(sh, [[(name, True, WHITE)]], 12, align=PP_ALIGN.LEFT, inset=8)
    for i, (spalte, tag) in enumerate(zeilen):
        yy = y + kopf_h + i * zeilen_h
        tb = slide.shapes.add_textbox(E(x + 8), E(yy + 1.5), E(w - 14), E(12.5))
        tf = tb.text_frame
        tf.word_wrap = False
        tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
        p = tf.paragraphs[0]
        r = p.add_run()
        r.text = spalte
        r.font.name = "Consolas"
        r.font.size = Pt(10.5)
        r.font.color.rgb = rgb(BODY)
        if tag:
            r2 = p.add_run()
            r2.text = "  " + tag
            r2.font.name = "Consolas"
            r2.font.size = Pt(10.5)
            r2.font.bold = True
            r2.font.color.rgb = rgb(METHOD)
    return y + h


# =====================================================================
# 1 — Deckblatt
# =====================================================================
s = folie("Frontpage_Digital")
for ph in s.placeholders:
    if ph.placeholder_format.idx == 10:
        ph.text_frame.text = ("Von der Miniwelt zur laufenden Datenbank — "
                              "eine Kontaktverwaltung mit SQLite")
    elif ph.placeholder_format.idx == 11:
        ph.text_frame.text = "Fallstudie mini-contacts"

# =====================================================================
# 2 — Kapitel 1
# =====================================================================
s = folie("Chapter")
s.shapes.title.text_frame.text = \
    "Kontaktverwaltung: der ganze Datenbankentwurf im Kleinen"

# =====================================================================
# 3 — Vorgehensmodell (Lisa, Chevron)
# =====================================================================
s = folie("Lisa_Slide")
kopf(s, "Fallstudie mini-contacts",
     "Sechs Schritte führen von der Miniwelt zur laufenden Anwendung",
     "Eigene Darstellung")
einl(s, "Datenbankentwurf ist ein Weg mit festen Etappen: Erst wird die "
        "Miniwelt beschrieben, dann entsteht das ER-Modell, die "
        "Normalisierung entfernt Redundanz, das relationale Modell übersetzt "
        "in Tabellen, SQLite setzt sie physisch um — erst am Ende steht die "
        "Anwendung. Wer die Reihenfolge einhält, baut Strukturen, die "
        "Änderungen überleben.")
chevron_kette(s, 176, 64, ["Miniwelt", "ER-Modell", "Normali-sierung",
                           "Relationales Modell", "SQLite", "CRUD-App"])
band(s, 300, 52, [[("Jede Etappe hat ein Ergebnis, das die nächste "
                    "braucht — Abkürzungen rächen sich spätestens beim "
                    "ersten Schemaumbau.", False, BODY)]])

# =====================================================================
# 4 — Miniwelt (Lisa, Kacheln)
# =====================================================================
s = folie("Lisa_Slide")
kopf(s, "Fallstudie mini-contacts",
     "Die Miniwelt legt fest, was die Datenbank abbilden muss",
     "Eigene Darstellung")
einl(s, "Die Miniwelt ist der Ausschnitt der Realität, den die Datenbank "
        "tragen soll — hier reichen vier Sätze, um sie zu beschreiben. Ihre "
        "Substantive werden später Entitäten, ihre Verben Beziehungen, und "
        "Mengenangaben wie \"beliebig viele\" legen die Kardinalitäten fest. "
        "Wer die Miniwelt sauber aufschreibt, hat den halben Entwurf.")
K4W = (CW - 3 * 14) / 4
for i, (titel, body) in enumerate([
        ("Kontakte", "Vorname, Nachname, optional Geburtstag und Notiz — "
                     "die zentrale Entität."),
        ("Kanäle", "Mobil, Festnetz, E-Mail, Web — beliebig viele je "
                   "Kontakt."),
        ("Adressen", "Privat, Arbeit, sonstige — keine, eine oder "
                     "mehrere je Kontakt."),
        ("Gruppen", "Familie, Arbeit, Verein — ein Kontakt in mehreren "
                    "Gruppen, eine Gruppe mit vielen Kontakten.")]):
    kachel(s, CL + i * (K4W + 14), 176, K4W, 150, BLUE, titel, [body])
band(s, 356, 52, [[("Miniwelt in vier Sätzen: ", True, BODY),
                   ("Substantive werden Entitäten, Verben werden "
                    "Beziehungen, Mengenangaben werden Kardinalitäten.",
                    False, BODY)]])

# =====================================================================
# 5 — Kapitel 2
# =====================================================================
s = folie("Chapter")
s.shapes.title.text_frame.text = \
    "Das ER-Modell beschreibt die Miniwelt vor jeder Tabelle"

# =====================================================================
# 6 — Chen-Bausteine (Lisa, Symbolreihe)
# =====================================================================
s = folie("Lisa_Slide")
kopf(s, "ER-Modell",
     "Chen-Notation trennt Entitäten, Beziehungen und Attribute",
     "Nach P. P. Chen (1976); eigene Darstellung")
einl(s, "Die Chen-Notation kennt genau drei Bausteine, und jeder hat eine "
        "feste Form: Rechtecke für Entitätstypen, Rauten für "
        "Beziehungstypen, Ellipsen für Attribute — Schlüsselattribute "
        "werden unterstrichen. Diese Trennung zwingt dazu, vor der ersten "
        "Tabelle zu klären, was ein Ding, was eine Verbindung und was nur "
        "eine Eigenschaft ist.")
symbole = [
    (MSO_SHAPE.RECTANGLE, "KONTAKT", 150, 48, False,
     "Rechteck — Entitätstyp",
     "Ein Ding der Miniwelt mit eigener Identität: Kontakt, Gruppe, "
     "Adresse."),
    (MSO_SHAPE.DIAMOND, "gehört zu", 170, 62, False,
     "Raute — Beziehungstyp",
     "Verbindet Entitätstypen; die Kardinalität steht an den Kanten: "
     "1:1, 1:N oder N:M."),
    (MSO_SHAPE.OVAL, "kontakt_id", 150, 44, True,
     "Ellipse — Attribut",
     "Eigenschaft einer Entität; das unterstrichene Schlüsselattribut "
     "identifiziert eindeutig."),
]
for i, (form, txt, w, h, unterstrich, titel, body) in enumerate(symbole):
    x = C3[i]
    sh = shape(s, form, x + (W3 - w) / 2 - 16, 176 + (66 - h) / 2, w, h,
               WHITE, BLUE, 1.25)
    label(sh, [[(txt, True, BLUE)]], 13)
    if unterstrich:
        for p in sh.text_frame.paragraphs:
            for r in p.runs:
                r.font.underline = True
    textbox(s, x, 262, W3 - 32, 18, [(titel, True, BLUE)], 13.5)
    textbox(s, x, 284, W3 - 32, 70, [body], 13)
band(s, 382, 52, [[("Erst die Miniwelt in diese drei Bausteine zerlegen — "
                    "Tabellen sind Schritt vier des Entwurfswegs, nicht "
                    "Schritt eins.", False, BODY)]])

# =====================================================================
# 7 — ER-Diagramm (Slide, Chen als Shapes)
# =====================================================================
s = folie("Slide")
kopf(s, "ER-Modell",
     "Ein Kontakt hat viele Kanäle und Adressen und gehört zu vielen Gruppen",
     "Eigene Darstellung")

# Attribute (Ellipsen) an KONTAKT — Auswahl, Rest im Repo
for txt, y, unterstrich in [("kontakt_id", 232, True), ("vorname", 292, False)]:
    sh = shape(s, MSO_SHAPE.OVAL, 90.8, y, 112, 32, WHITE, HAIR, 1.0)
    label(sh, [[(txt, False, BODY)]], 12)
    if unterstrich:
        for p in sh.text_frame.paragraphs:
            for r in p.runs:
                r.font.underline = True
# Kanten Attribut -> Entitaet
hline(s, 202.8, 248, 22)
vline(s, 224.8, 248, 57)
hline(s, 224.8, 305, 25.2)
hline(s, 202.8, 308, 47.2)

# Kernentitaet
sh = shape(s, MSO_SHAPE.RECTANGLE, 250, 280, 150, 52, BLUE)
label(sh, [[("KONTAKT", True, WHITE)]], 14)

# Rauten
for txt, y in [("erreichbar über", 186), ("wohnt an", 276),
               ("gehört zu", 374)]:
    sh = shape(s, MSO_SHAPE.DIAMOND, 470, y, 132, 60, WHITE, BLUE, 1.25)
    label(sh, [[(txt, False, BLUE)]], 10.5, inset=1)

# Entitaeten rechts
for txt, y in [("KANAL", 191), ("ADRESSE", 281), ("GRUPPE", 379)]:
    sh = shape(s, MSO_SHAPE.RECTANGLE, 690, y, 150, 50, WHITE, BLUE, 1.25)
    label(sh, [[(txt, True, BLUE)]], 13)

# Kanten KONTAKT -> Rauten (orthogonal)
hline(s, 400, 292, 35)          # Ausgang oben -> R1
vline(s, 435, 216, 76)
hline(s, 435, 216, 35)
hline(s, 400, 306, 70)          # Mitte -> R2
hline(s, 400, 320, 35)          # Ausgang unten -> R3
vline(s, 435, 320, 84)
hline(s, 435, 404, 35)
# Kanten Rauten -> Entitaeten
hline(s, 602, 216, 88)
hline(s, 602, 306, 88)
hline(s, 602, 404, 88)
# Kardinalitaeten
kardlabel(s, 440, 197, "1")
kardlabel(s, 404, 289, "1")
kardlabel(s, 440, 385, "M")
kardlabel(s, 664, 199, "N")
kardlabel(s, 664, 289, "N")
kardlabel(s, 664, 387, "N")

textbox(s, CL, 462, CW, 18,
        [("Schlüsselattribute unterstrichen; die übrigen Attribute je "
          "Entität stehen im Repo (docs/02_er_modell.md).", False, SEC)],
        11.5)

# =====================================================================
# 8 — Kapitel 3
# =====================================================================
s = folie("Chapter")
s.shapes.title.text_frame.text = \
    "Normalisierung entfernt Redundanz Schritt für Schritt"

# =====================================================================
# 9 — Flache Liste und Anomalien (Lisa)
# =====================================================================
s = folie("Lisa_Slide")
kopf(s, "Normalisierung",
     "Die flache Kontaktliste erzwingt Redundanz und Anomalien",
     "Eigene Darstellung")
einl(s, "Als einzelne Tabelle gedacht, mischt die Kontaktliste mehrere "
        "Werte in eine Zelle und wiederholt denselben Namen Zeile für "
        "Zeile. Solche Strukturen erzeugen drei typische Anomalien: Ändern "
        "trifft nie alle Kopien, Löschen nimmt Wissen mit, und Einfügen "
        "braucht Zeilen, die es noch gar nicht gibt.")
SP = [(90.8, 168), (258.8, 234), (492.8, 262), (754.8, 238)]
zellen = [
    ("name", "telefon", "adresse", "gruppen"),
    ("Anna Müller", "+49 171 111, +49 931 222",
     "Hauptstr. 5, 97070 Würzburg", "Familie, Verein"),
    ("Ben Bauer", "+49 172 333", "Mainkai 53, 97082 Würzburg", "Arbeit"),
    ("Anna Müller", "+49 171 111", "Gartenweg 2, 97318 Kitzingen",
     "Familie"),
]
rot = {(1, 1), (1, 3), (3, 0)}                 # Listen in Zellen, Dublette
for zi, zeile in enumerate(zellen):
    y = 176 + (0 if zi == 0 else 24 + (zi - 1) * 26)
    h = 24 if zi == 0 else 26
    for si, txt in enumerate(zeile):
        x, w = SP[si]
        sh = shape(s, MSO_SHAPE.RECTANGLE, x, y, w, h,
                   SAND if zi == 0 else WHITE, HAIR)
        farbe = WARN if (zi, si) in rot else BODY
        label(sh, [[(txt, zi == 0, BLUE if zi == 0 else farbe)]], 11.5,
              align=PP_ALIGN.LEFT, inset=8)
for i, (titel, body) in enumerate([
        ("Änderungsanomalie", "Annas Nummer steht in mehreren Zeilen — "
         "eine vergessene Kopie macht den Bestand widersprüchlich."),
        ("Löschanomalie", "Mit Bens letzter Zeile verschwindet auch das "
         "Wissen, dass es die Gruppe \"Arbeit\" gibt."),
        ("Einfügeanomalie", "Eine neue Gruppe ohne Mitglied lässt sich "
         "gar nicht erfassen — es gibt keine Zeile dafür.")]):
    kachel(s, C3[i], 316, W3, 138, WARN, titel, [body])

# =====================================================================
# 10 — 1NF (Lisa, Gegenueberstellung)
# =====================================================================
s = folie("Lisa_Slide")
kopf(s, "Normalisierung",
     "Die erste Normalform macht jeden Attributwert atomar",
     "Eigene Darstellung")
einl(s, "Die erste Normalform verlangt genau einen Wert je Zelle. Listen "
        "wie \"Familie, Verein\" werden zu einzelnen Zeilen, "
        "zusammengesetzte Angaben wie die Adresse in ihre Bestandteile "
        "zerlegt. Die Daten werden dadurch erst durchsuchbar und "
        "sortierbar — die Redundanz aber wächst zunächst sogar noch an.")
gegenueber(
    s, 176, 240,
    (WARN, "Vorher: Listen in der Zelle",
     [[("telefon = ", False, BODY),
       ("\"+49 171 111, +49 931 222\"", False, WARN)],
      [("gruppen = ", False, BODY), ("\"Familie, Verein\"", False, WARN)],
      "Eine Zeile je Kontakt — nicht durchsuchbar, nicht sortierbar, "
      "nicht einzeln löschbar."]),
    (GOOD, "Nachher: eine Zeile je Wert",
     [[("kanal(1, mobil, \"+49 171 111\")", False, BODY)],
      [("kanal(1, festnetz, \"+49 931 222\")", False, BODY)],
      "Jeder Wert atomar in einer eigenen Zeile — Struktur sauber, "
      "Redundanz beseitigen erst 2NF und 3NF."]),
    badge_l="✗", badge_r="✓")
band(s, 440, 44, [[("Merksatz: ", True, BODY),
                   ("1NF ordnet die Struktur — gegen Redundanz helfen "
                    "erst die nächsten Normalformen.", False, BODY)]])

# =====================================================================
# 11 — 2NF/3NF (Lisa, Zuordnung)
# =====================================================================
s = folie("Lisa_Slide")
kopf(s, "Normalisierung",
     "2NF und 3NF lösen versteckte Abhängigkeiten auf",
     "Eigene Darstellung")
einl(s, "Nach der 1NF hängen noch Attribute am falschen Ort: Der Name "
        "hängt nur an der Kontaktnummer statt am ganzen zusammengesetzten "
        "Schlüssel, und der Ort hängt über die Postleitzahl nur mittelbar "
        "daran. 2NF und 3NF zerlegen die Tabelle entlang dieser "
        "Abhängigkeiten in eigene Relationen.")
zuordnung(s, 176, [
    ("2NF — ganzer Schlüssel",
     "Name und Geburtstag hängen nur an kontakt_id — sie wandern in die "
     "eigene Tabelle kontakt."),
    ("3NF — nichts Transitives",
     "plz → ort: Der Ort hängt am Schlüssel nur über die Postleitzahl — "
     "streng genommen gehört er in eine Ort-Tabelle."),
    ("Pragmatik — bewusst abweichen",
     "Wir behalten plz und ort zusammen: dokumentierte Entscheidung im "
     "Schema statt blinder Regeltreue."),
], bw=232)
band(s, 350, 52, [[("Normalisierung ist ein Werkzeug, um Redundanz zu "
                    "erkennen — nicht jede erkannte Redundanz muss "
                    "beseitigt werden.", False, BODY)]])

# =====================================================================
# 12 — Kapitel 4
# =====================================================================
s = folie("Chapter")
s.shapes.title.text_frame.text = "Vom ER-Modell zum relationalen Modell"

# =====================================================================
# 13 — Abbildungsregeln (Lisa, Kacheln)
# =====================================================================
s = folie("Lisa_Slide")
kopf(s, "Relationales Modell",
     "Drei Regeln übersetzen das ER-Modell in Tabellen",
     "Eigene Darstellung")
einl(s, "Die Übersetzung ins relationale Modell folgt festen Regeln und "
        "lässt keinen Spielraum: Entitätstypen werden Tabellen, "
        "1:N-Beziehungen werden zum Fremdschlüssel auf der N-Seite, und "
        "N:M-Beziehungen brauchen eine eigene Koppeltabelle, weil sie "
        "sich in keiner einzelnen Spalte unterbringen lassen.")
for i, (titel, body) in enumerate([
        ("Entität → Tabelle", "KONTAKT wird zur Tabelle kontakt; das "
         "Schlüsselattribut wird Primärschlüssel kontakt_id."),
        ("1:N → Fremdschlüssel", "kanal.kontakt_id verweist auf kontakt — "
         "die N-Seite trägt den Schlüssel der 1-Seite."),
        ("N:M → Koppeltabelle", "kontakt_gruppe(kontakt_id, gruppe_id): "
         "zwei Fremdschlüssel, zusammen der Primärschlüssel.")]):
    kachel(s, C3[i], 176, W3, 150, METHOD, titel, [body])
band(s, 356, 44, [[("Merkregel: ", True, BODY),
                   ("Beziehungen verschwinden nicht — sie werden zu "
                    "Schlüsseln.", False, BODY)]])

# =====================================================================
# 14 — Relationales Schema (Slide, Tabellenboxen)
# =====================================================================
s = folie("Slide")
kopf(s, "Relationales Modell",
     "Fünf Tabellen tragen das gesamte Datenmodell",
     "Eigene Darstellung")
tabellenbox(s, 90.8, 236, 190, "kontakt", [
    ("kontakt_id", "PK"), ("vorname", ""), ("nachname", ""),
    ("geburtstag", ""), ("notiz", ""), ("angelegt_am", "")])
tabellenbox(s, 400, 176, 190, "kanal", [
    ("kanal_id", "PK"), ("kontakt_id", "FK"), ("typ", ""), ("wert", "")])
tabellenbox(s, 400, 286, 190, "adresse", [
    ("adresse_id", "PK"), ("kontakt_id", "FK"), ("typ", ""),
    ("strasse", ""), ("hausnummer", ""), ("plz", ""), ("ort", "")])
tabellenbox(s, 710, 176, 190, "gruppe", [
    ("gruppe_id", "PK"), ("name", ""), ("farbe", "")])
tabellenbox(s, 650, 396, 220, "kontakt_gruppe", [
    ("kontakt_id", "PK FK"), ("gruppe_id", "PK FK")])
# Kanten (orthogonal): N-Seite -> 1-Seite
hline(s, 280.8, 250, 69.2)      # kontakt -> kanal (Eingang y=217)
vline(s, 350, 217, 33)
hline(s, 350, 217, 50)
hline(s, 280.8, 320, 69.2)      # kontakt -> adresse (Eingang y=349)
vline(s, 350, 320, 29)
hline(s, 350, 349, 50)
vline(s, 805, 243, 153)         # gruppe -> kontakt_gruppe
vline(s, 185, 348, 74)          # kontakt -> kontakt_gruppe
hline(s, 185, 422, 465)
kardlabel(s, 286, 232, "1")
kardlabel(s, 374, 199, "N")
kardlabel(s, 286, 324, "1")
kardlabel(s, 374, 331, "N")
kardlabel(s, 810, 248, "1")
kardlabel(s, 810, 372, "N")
kardlabel(s, 190, 352, "1")
kardlabel(s, 262, 404, "N")
textbox(s, 90.8, 462, CW, 18,
        [("Die N:M-Beziehung aus dem ER-Modell lebt jetzt in der "
          "Koppeltabelle — zu beiden Seiten je eine 1:N-Beziehung.",
          False, SEC)], 11.5)

# =====================================================================
# 15 — Kapitel 5
# =====================================================================
s = folie("Chapter")
s.shapes.title.text_frame.text = "SQLite: die Datenbank ohne Server"

# =====================================================================
# 16 — SQLite vs. Server (Lisa, Gegenueberstellung)
# =====================================================================
s = folie("Lisa_Slide")
kopf(s, "SQLite",
     "SQLite ist eine Datei im Prozess der Anwendung — kein Server",
     "sqlite.org/mostdeployed.html (Stand 08/2026); eigene Darstellung")
einl(s, "Klassische Datenbanksysteme laufen als eigener Serverprozess mit "
        "Netzwerk, Konten und Administration. SQLite verfolgt das "
        "Gegenmodell: Die komplette Datenbank ist eine einzige Datei, die "
        "Bibliothek läuft im Prozess der Anwendung — SQL, Schlüssel und "
        "Normalformen funktionieren dabei identisch zu den großen "
        "Systemen.")
gegenueber(
    s, 176, 216,
    (BLUE, "Client-Server (PostgreSQL, MySQL)",
     [["eigener Serverprozess, Zugriff über das Netzwerk"],
      ["Benutzerverwaltung und Rechte"],
      ["viele gleichzeitige Schreiber — zentrale Unternehmensdaten"]]),
    (METHOD, "SQLite (eingebettet)",
     [["eine Datei: contacts.db — kopieren ist Backup"],
      ["in Python eingebaut: import sqlite3"],
      ["ein Schreiber zur Zeit — lokale Apps, Mobile, Browser"]]),
    badge_l="▸", badge_r="▸")
band(s, 416, 52, [[("SQLite steckt in jedem Smartphone und Browser — die "
                    "am weitesten verbreitete Datenbank der Welt. Für die "
                    "Lehre heißt das: gleiches SQL, null Aufbau.",
                    False, BODY)]])

# =====================================================================
# 17 — DDL (Tool, Code + Sprechblasen)
# =====================================================================
s = folie("Tool_Slide")
kopf(s, "SQLite",
     "CREATE TABLE setzt das Modell wörtlich in SQL um",
     "mini-contacts, db/schema.sql")
einl(s, "Das physische Schema ist die wörtliche Übersetzung des "
        "relationalen Modells: je Tabelle ein CREATE TABLE, je Beziehung "
        "ein Fremdschlüssel. Regeln, die immer gelten sollen, gehören in "
        "das Schema — Wertebereiche als CHECK, Eindeutigkeit als UNIQUE, "
        "Löschverhalten als ON DELETE CASCADE.", variante="Tool_Slide")
shape(s, MSO_SHAPE.RECTANGLE, 90.8, 176, 540, 232, SAND)
mono_box(s, 108, 192, 506, 200, [
    "CREATE TABLE kanal (",
    "    kanal_id    INTEGER PRIMARY KEY,",
    "    kontakt_id  INTEGER NOT NULL",
    "                REFERENCES kontakt(kontakt_id)",
    "                ON DELETE CASCADE,",
    "    typ   TEXT NOT NULL",
    "          CHECK (typ IN ('mobil', 'festnetz',",
    "                         'email', 'web')),",
    "    wert  TEXT NOT NULL",
    ");",
])
sprechblase(s, 660, 200, 330, 84,
            ["ON DELETE CASCADE: Kanäle sterben mit ihrem Kontakt — die "
             "Datenbank räumt auf, nicht die App."], farbe=METHOD, nummer=1)
sprechblase(s, 660, 306, 330, 84,
            ["CHECK erzwingt den Wertebereich direkt im Schema — ein "
             "Enum auf Datenbankebene."], farbe=METHOD, nummer=2)

# =====================================================================
# 18 — PRAGMA foreign_keys (Tool, Kacheln + Codeband)
# =====================================================================
s = folie("Tool_Slide")
kopf(s, "SQLite",
     "Ohne PRAGMA prüft SQLite keine Fremdschlüssel",
     "sqlite.org/foreignkeys.html (Stand 08/2026)")
einl(s, "Aus Kompatibilität zu alten Versionen ist die "
        "Fremdschlüsselprüfung in SQLite standardmäßig abgeschaltet — "
        "REFERENCES ist dann nur Dokumentation. Eingeschaltet wird die "
        "Prüfung je Verbindung; der Aufruf gehört deshalb an genau eine "
        "Stelle: direkt neben das Öffnen der Datenbank.",
     variante="Tool_Slide")
gegenueber(
    s, 176, 190,
    (WARN, "Pragma vergessen: stille Fehler",
     ["INSERT mit kontakt_id 9999 gelingt, DELETE hinterlässt verwaiste "
      "Kanäle — kein Fehler, aber falsche Daten."]),
    (GOOD, "Pragma gesetzt: laute Fehler",
     ["\"FOREIGN KEY constraint failed\" — der Fehler kommt sofort und "
      "zeigt auf die Ursache statt auf Symptome."]),
    badge_l="✗", badge_r="✓")
shape(s, MSO_SHAPE.RECTANGLE, CL, 396, CW, 60, SAND)
mono_box(s, CL + 18, 410, CW - 36, 34, [
    "con = sqlite3.connect(\"contacts.db\")",
    "con.execute(\"PRAGMA foreign_keys = ON\")   # bei jedem Connect",
])

# =====================================================================
# 19 — Kapitel 6
# =====================================================================
s = folie("Chapter")
s.shapes.title.text_frame.text = "Eine kleine Web-App macht CRUD greifbar"

# =====================================================================
# 20 — Architektur (Slide, Fluss + Kennzahlen)
# =====================================================================
s = folie("Slide")
kopf(s, "Anwendung",
     "Jeder Klick in der Oberfläche wird zu genau einem SQL-Befehl",
     "Eigene Darstellung")
for txt, sub, x, fill in [("Browser", "HTML + CSS", 90.8, None),
                          ("Flask", "app.py — Routen", 412, None),
                          ("SQLite", "contacts.db", 733.2, BLUE)]:
    sh = shape(s, MSO_SHAPE.RECTANGLE, x, 176, 200, 64,
               fill or WHITE, None if fill else HAIR, 1.0)
    label(sh, [[(txt, True, WHITE if fill else BLUE)],
               [(sub, False, WHITE if fill else SEC)]], 13)
for x in (302, 623.2):
    shape(s, MSO_SHAPE.RIGHT_ARROW, x + 8, 200, 94, 16, BLUE)
textbox(s, 250, 250, 216, 16,
        [[("HTTP POST — ", True, SEC), ("/kontakt/24/loeschen", False, SEC)]],
        11, align=PP_ALIGN.CENTER)
textbox(s, 566, 250, 216, 16,
        [[("SQL — ", True, SEC), ("DELETE FROM kontakt …", False, SEC)]],
        11, align=PP_ALIGN.CENTER)
kennzahlen(s, 310, 92, [("Create", "INSERT", BLUE), ("Read", "SELECT", BLUE),
                        ("Update", "UPDATE", BLUE),
                        ("Delete", "DELETE", BLUE)])
textbox(s, CL, 430, CW, 18,
        [("Vier Grundoperationen — mehr braucht eine datengetriebene "
          "Anwendung nicht, und jede hat ihr SQL-Gegenstück.", False, SEC)],
        11.5)

# =====================================================================
# 21 — SQL-Log (Tool, Screenshot + Sprechblasen)
# =====================================================================
s = folie("Tool_Slide")
kopf(s, "Anwendung",
     "Die App zeigt zu jedem Klick das ausgeführte SQL",
     "Eigene Darstellung (mini-contacts, Stand 08/2026)")
einl(s, "Das didaktische Herzstück der App ist das SQL-Log am unteren "
        "Rand: Es zeigt die zuletzt ausgeführten Statements mit "
        "eingesetzten Parametern. Suche, Gruppenfilter und jedes Speichern "
        "lassen sich so live auf ihr SQL zurückführen — die Oberfläche "
        "wird zur Übungsumgebung für die Vorlesung.",
     variante="Tool_Slide")
s.shapes.add_picture(str(IMG / "ui_uebersicht.png"), E(90.8), E(176),
                     height=E(300))
sprechblase(s, 630, 210, 360, 80,
            ["Gruppen-Chips filtern die Liste — dahinter liegt ein JOIN "
             "über die Koppeltabelle kontakt_gruppe."],
            farbe=METHOD, nummer=1)
sprechblase(s, 630, 320, 360, 80,
            ["SQL-Log am unteren Rand: die letzten Statements, das "
             "jüngste zuoberst."], farbe=METHOD, nummer=2)

# =====================================================================
# 22 — Fake-Daten (Bot, Kacheln)
# =====================================================================
s = folie("Bot_Slide")
kopf(s, "Anwendung",
     "Fake-Daten kommen aus Generatoren — nie aus echten Adressbüchern",
     "mockaroo.com; faker.readthedocs.io (Stand 08/2026)")
einl(s, "Eine leere Datenbank lässt sich weder testen noch vorführen, "
        "echte Personendaten sind tabu — Testdaten werden deshalb "
        "generiert. Drei Wege führen zum Ziel: konfigurierbare "
        "Generatoren im Browser, Programmbibliotheken im eigenen Skript "
        "und KI-Assistenten, die das Schema samt Constraints verstehen.",
     variante="Bot_Slide")
for i, (titel, body) in enumerate([
        ("Mockaroo (Browser)", "Spalten und Datentypen zusammenklicken, "
         "bis 1 000 Zeilen gratis — Export als CSV, JSON oder direkt als "
         "SQL-INSERTs."),
        ("Faker (Python)", "pip install faker, Locale de_DE — Namen, "
         "Adressen und Mails in beliebiger Menge, reproduzierbar per "
         "Seed."),
        ("KI-Assistenten", "Schema hingeben, INSERTs generieren lassen — "
         "versteht CHECK und Fremdschlüssel; Ergebnis immer gegen die "
         "Datenbank laufen lassen.")]):
    kachel(s, C3[i], 176, W3, 168, METHOD, titel, [body])
band(s, 372, 52, [[("Datenschutz: ", True, BODY),
                   ("Fake-Daten sind Pflicht, sobald ein Repo öffentlich "
                    "wird — E-Mail-Domains nach RFC 2606 (example.org).",
                    False, BODY)]])

# =====================================================================
# 23 — Repo (Slide, zwei Karten)
# =====================================================================
s = folie("Slide")
kopf(s, "Anwendung",
     "Das Repo mini-contacts enthält den ganzen Weg zum Nachbauen",
     "github.com/swrobuts/mini-contacts")
shape(s, MSO_SHAPE.RECTANGLE, 90.8, 176, 430, 246, SAND)
textbox(s, 108, 188, 396, 18, [("Struktur", True, BLUE)], 13)
mono_box(s, 108, 212, 396, 198, [
    "mini-contacts/",
    "├── app.py           # Flask-App + SQL-Log",
    "├── seed.py          # Fake-Daten-Generator",
    "├── db/schema.sql    # das physische Schema",
    "├── docs/            # Entwurfsweg Schritt 1-7",
    "│                    #   mit Mermaid-Diagrammen",
    "├── slides/          # dieses Foliendeck",
    "├── static/          # Styling",
    "└── templates/       # HTML (Jinja2)",
], size=12)
shape(s, MSO_SHAPE.RECTANGLE, 560, 176, 430, 246, SAND)
textbox(s, 578, 188, 396, 18, [("Schnellstart (Win / macOS / Linux)",
                                True, BLUE)], 13)
mono_box(s, 578, 212, 396, 198, [
    "git clone github.com/swrobuts/",
    "          mini-contacts.git",
    "cd mini-contacts",
    "pip install flask",
    "python seed.py    # 48 Demo-Kontakte",
    "python app.py",
    "# -> http://127.0.0.1:5050",
], size=12)
textbox(s, 90.8, 446, CW, 22,
        [[("github.com/swrobuts/mini-contacts", True, BLUE),
          ("  — Fragen und Fundstücke gern als Issue.", False, SEC)]], 14)

# =====================================================================
prs.save(str(ZIEL))
print(f"Gespeichert: {ZIEL} ({len(prs.slides._sldIdLst)} Folien)")
B.report_warnings()
